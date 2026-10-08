"""Add or replace the sheet "MASQ replication" in outputs/all_runs/all_runs_results.xlsx (standard 40-column run layout).

Base python (openpyxl), no torch. Reads a finished run's manifest and units/*/scores.json only. One row per (arm, dataset),
seed means. Configurations and scores only: no comparison columns (the sheet is a record of configurations for reuse).
Back up the workbook first; on PermissionError (Excel open) a staged copy is written instead and nothing is forced.

  python -m script.week5.masq_experiments.sheet --run-id masq_v1 [--workbook PATH] [--out PATH]

A smoke run is never written to the real workbook (pass --out to write a staged test copy).
"""
from __future__ import annotations

import argparse
import copy
import json
import shutil
import statistics as st
from pathlib import Path

import openpyxl

from script.week5.masq_experiments import config as C

SHEET = "MASQ replication"
REAL = C.ROOT / "outputs" / "all_runs" / "all_runs_results.xlsx"
ARM_ROLE = {"S0": "SMQ retrained: the paired control", "M-pub": "MASQ, published settings (reimplementation from the paper text)",
            "M-JL": "JLSD only (no velocity loss)", "M-fix": "MASQ with one label-free dropout rate p=0.25 on every dataset",
            "V1": "JLSD + coordinate velocity loss with a scale-free calibrated weight", "V2a": "JLSD + latent smoothness prior (rho 0.01)",
            "V2b": "JLSD + latent smoothness prior (rho 0.1)"}
COUNT_NAMES = {"hugadb": "HuGaDB", "lara": "LARa", "babel1": "BABEL-1", "babel2": "BABEL-2", "babel3": "BABEL-3"}


def loss_text(arm: str, p: float) -> str:
    base = "joint-distance reconstruction + commit"
    if arm == "S0":
        return base
    if arm == "M-JL":
        return f"{base} + JLSD p={p:g}"
    if arm == "V1":
        return f"{base} + JLSD p={p:g} + mask-aware velocity (scale-free calibrated weight, rho={C.VARIANT_RHO[arm]:g})"
    if arm in C.VARIANT_ARMS:
        return f"{base} + JLSD p={p:g} + latent smoothness prior (calibrated, rho={C.VARIANT_RHO[arm]:g})"
    return f"{base} + JLSD p={p:g} + mask-aware velocity (lambda={C.VEL_WEIGHT:g})"


def build_rows(run_dir: Path, smoke: bool) -> list[list]:
    rid = run_dir.name
    cells = {}
    for sp in sorted((run_dir / "units").glob("*/scores.json")):
        u = C.parse_unit(sp.parent.name)
        cells.setdefault((u.arm, u.dataset), {})[u.seed] = json.loads(sp.read_text())
    rows = []
    for arm in C.ARMS:
        for ds in ("hugadb", "lara", "babel1", "babel2", "babel3"):
            seeds = cells.get((arm, ds))
            if not seeds:
                continue
            ss = [s for s in C.SEEDS if s in seeds]
            u = C.Unit(arm, ds, ss[0])
            mean = {m: round(st.mean(seeds[s]["global"][m] for s in ss), 2) for m in C.METRICS}
            segs = st.mean(seeds[s]["segments"]["mean_pred"] for s in ss)
            gts = st.mean(seeds[s]["segments"]["mean_gt"] for s in ss)
            note = (f"{arm}: {ARM_ROLE[arm]}. Final epoch ({C.SHARED['epochs']}), global Hungarian, pooled/transductive; seed-mean shown; "
                    f"K = C = {seeds[ss[0]]['codes']['num_codes']} codes. Predicted {segs:.1f} vs ground-truth {gts:.1f} segments per recording."
                    + (f" Incomplete: {len(ss)} of {len(C.SEEDS)} seeds." if len(ss) < len(C.SEEDS) else "")
                    + (" SMOKE RUN: not a result." if smoke else "") + " Development result (protocol P7).")
            rows.append([f"{rid} ({arm})", f"results\\week5\\masq_experiments\\{rid}\\REPORT.md", COUNT_NAMES[ds], ", ".join(map(str, ss)),
                         "TCN joint-disentangled (SMQ backbone), trained", 29, "vq (EMA)", "k=c", loss_text(arm, u.jlsd_p),
                         "na", "na", "na", "na", "na",  # preprocessing .. discretization method, as the e0 rows
                         seeds[ss[0]]["codes"]["num_codes"], C.DATASETS[ds]["patch_size"],
                         "none", *["na"] * 11, "none", *["na"] * 5, *[mean[m] for m in C.METRICS], note])
    return rows


def write_sheet(workbook: Path, out: Path, rows: list[list]) -> Path:
    wb = openpyxl.load_workbook(workbook)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    runs = wb["Runs"]
    for r in range(1, 4):
        for c in range(1, 41):
            a, b = runs.cell(r, c), ws.cell(r, c)
            b.value = a.value
            if a.has_style:
                b.font, b.border, b.fill = copy.copy(a.font), copy.copy(a.border), copy.copy(a.fill)
                b.number_format, b.alignment, b.protection = a.number_format, copy.copy(a.alignment), copy.copy(a.protection)
    for rng in runs.merged_cells.ranges:
        if rng.max_row <= 3:
            ws.merge_cells(str(rng))
    for key, dim in runs.column_dimensions.items():
        ws.column_dimensions[key].width = dim.width
    for key, dim in runs.row_dimensions.items():
        if key <= 3 and dim.height:
            ws.row_dimensions[key].height = dim.height
    for i, row in enumerate(rows, 4):
        assert len(row) == 40, len(row)
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
    try:
        wb.save(out)
        return out
    except PermissionError:
        staged = out.with_name(out.stem + ".staged.xlsx")
        wb.save(staged)
        print(f"PermissionError on {out}: Excel probably has it open. Wrote {staged} instead; close Excel and re-run.")
        return staged


def main(argv=None) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--run-id", required=True)
    ap.add_argument("--workbook", type=Path, default=REAL)
    ap.add_argument("--out", type=Path, help="default: the workbook itself (refused for smoke runs)")
    args = ap.parse_args(argv)
    run_dir = C.run_dir(args.run_id)
    manifest = json.loads((run_dir / "manifest.json").read_text())
    smoke = bool(manifest.get("smoke"))
    out = args.out or args.workbook
    if smoke and out.resolve() == REAL.resolve():
        raise SystemExit("refusing to write a smoke run into the real workbook; pass --out for a staged test copy")
    rows = build_rows(run_dir, smoke)
    if not rows:
        raise SystemExit("no scored units in this run")
    if out.resolve() == args.workbook.resolve() and out.exists():
        backup = run_dir / "all_runs_results.backup.xlsx"
        shutil.copy2(out, backup)
        print(f"backup: {backup}")
    print(f"wrote {write_sheet(args.workbook, out, rows)}: {len(rows)} rows")


if __name__ == "__main__":
    main()
